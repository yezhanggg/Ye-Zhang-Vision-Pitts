"""Read each public source into a tract-level table for the city tracts, keyed by the 11-digit 2020 GEOID.

Every function returns a DataFrame indexed by GEOID with raw (unranked) columns. Ranking, composites and
confidence tags happen in factors.py. Nothing here imputes a value: missing stays missing.
"""
from __future__ import annotations

import json
import zipfile

import geopandas as gpd
import numpy as np
import openpyxl
import pandas as pd
import requests

from visionpitts.config import (
    ACS_YEAR,
    CENSUS_API_KEY,
    COUNTY_FIPS,
    COUNTY_GEOID,
    CRS_PA_SOUTH,
    CRS_WGS84,
    PROCESSED,
    RAW,
    SQM_PER_ACRE,
    STATE_FIPS,
)
from visionpitts import geo

Z90 = 1.645  # ACS margins of error are published at the 90% level


def cv(est: pd.Series, moe: pd.Series) -> pd.Series:
    """Coefficient of variation from a 90% MOE. NaN where the estimate is 0 or missing."""
    e = est.where(est > 0)
    return (moe / Z90) / e


# ----------------------------------------------------------------------------------------------- ACS
ACS_VARS = {
    "B01003_001": "pop",
    "B11001_001": "households",
    "B19013_001": "med_hh_income",
    "B25064_001": "med_gross_rent",
    "B25077_001": "med_home_value",
    "B25003_001": "occ_units",
    "B25003_003": "renter_hh",
    "B25070_001": "rent_hh_total",
    "B25070_007": "rb30",
    "B25070_008": "rb35",
    "B25070_009": "rb40",
    "B25070_010": "rb50",
    "B25002_001": "units_total",
    "B25002_003": "units_vacant",
}


def acs(geoids: pd.Index) -> pd.DataFrame:
    """ACS 5-year estimates + MOE from the Census API (falls back to the cached raw CSV if the API is unreachable)."""
    cols = [f"{v}{s}" for v in ACS_VARS for s in ("E", "M")]
    url = f"https://api.census.gov/data/{ACS_YEAR}/acs/acs5"
    params = {"get": ",".join(cols), "for": "tract:*", "in": f"state:{STATE_FIPS} county:{COUNTY_FIPS}"}
    if CENSUS_API_KEY:
        params["key"] = CENSUS_API_KEY
    try:
        r = requests.get(url, params=params, timeout=60)
        r.raise_for_status()
        rows = r.json()
        df = pd.DataFrame(rows[1:], columns=rows[0])
        df["GEOID"] = df["state"] + df["county"] + df["tract"]
        print(f"  ACS: {len(df)} county tracts from the Census API")
    except Exception as e:  # offline fallback
        print(f"  ACS API failed ({e}); using cached download")
        df = pd.read_csv(RAW / "acs2024_tracts_raw.csv", dtype={"GEOID": str})
    df = df.set_index("GEOID")
    for c in cols:
        df[c] = pd.to_numeric(df[c], errors="coerce")
        if c.endswith("E"):
            df.loc[df[c] < 0, c] = np.nan            # -666666666 and friends: not available
        else:
            df.loc[df[c] == -555555555, c] = 0.0     # controlled estimate: MOE not applicable
            df.loc[df[c] < 0, c] = np.nan
    out = pd.DataFrame(index=df.index)
    for var, name in ACS_VARS.items():
        out[name] = df[f"{var}E"]
        out[f"{name}_moe"] = df[f"{var}M"]
    out["med_hh_income_cv"] = cv(out["med_hh_income"], out["med_hh_income_moe"])
    out["med_gross_rent_cv"] = cv(out["med_gross_rent"], out["med_gross_rent_moe"])
    out["renter_share"] = out["renter_hh"] / out["occ_units"].where(out["occ_units"] > 0)
    burdened = out[["rb30", "rb35", "rb40", "rb50"]].sum(axis=1, min_count=1)
    out["rent_burdened_share"] = burdened / out["rent_hh_total"].where(out["rent_hh_total"] > 0)
    out["vacancy_share"] = out["units_vacant"] / out["units_total"].where(out["units_total"] > 0)
    keep = ["pop", "households", "med_hh_income", "med_hh_income_cv", "med_gross_rent", "med_gross_rent_cv",
            "med_home_value", "renter_hh", "renter_hh_moe", "renter_share", "rent_burdened_share", "vacancy_share"]
    return out[keep].reindex(geoids)


# ---------------------------------------------------------------------------------------- ACS shares
ACS_TRACT_CSV = PROCESSED / "acs_tract.csv"  # written by scripts/07_build_acs_levels.py
ACS_SHARE_COLS = ["age65_share", "age65_share_cv", "units_2_4_share", "units_2_4_share_cv"]
ACS_SHARE_OPTIONAL = ["households_cv"]


def acs_shares(geoids: pd.Index, path=None) -> pd.DataFrame:
    """Two ACS shares that are scored (residents 65+, units in 2-4 unit buildings) with their CVs, plus the
    household-count CV when the table has it. Read from data/processed/acs_tract.csv; nothing is imputed."""
    path = ACS_TRACT_CSV if path is None else path
    if not path.exists():
        raise FileNotFoundError(f"{path} is missing. Run scripts/07_build_acs_levels.py before scripts/02_build_factors.py "
                                "(pipeline order: 01, 07, 02, 03, 04).")
    a = pd.read_csv(path, dtype={"GEOID": str})
    missing = [c for c in ["GEOID", *ACS_SHARE_COLS] if c not in a.columns]
    if missing:
        raise KeyError(f"{path} has no column {missing}. Rebuild it with scripts/07_build_acs_levels.py; the scored ACS "
                       f"shares need {ACS_SHARE_COLS}.")
    a = a.set_index("GEOID")
    keep = ACS_SHARE_COLS + [c for c in ACS_SHARE_OPTIONAL if c in a.columns]
    out = a[keep].apply(pd.to_numeric, errors="coerce").reindex(geoids)
    print(f"  ACS shares: {int(out['age65_share'].notna().sum())} tracts with a 65+ share, "
          f"{int(out['units_2_4_share'].notna().sum())} with a 2-4 unit share")
    return out


# ---------------------------------------------------------------------------------------------- CHAS
CHAS_ZIP = RAW / "benchmark" / "chas" / "2018thru2022-140-csv.zip"
CHAS_DICT = RAW / "benchmark" / "chas" / "CHAS-data-dictionary-18-22.xlsx"
LOW_INCOME = ["less than or equal to 30% of HAMFI", "greater than 30% but less than or equal to 50% of HAMFI"]
BURDENED = ["greater than 30% but less than or equal to 50%", "greater than 50%"]


def chas_columns() -> tuple[list[str], list[str]]:
    """Pick Table 8 columns by their dictionary description rather than by hard-coded names."""
    ws = openpyxl.load_workbook(CHAS_DICT, read_only=True)["Table 8"]
    rows = list(ws.iter_rows(values_only=True))
    d = pd.DataFrame(rows[1:], columns=["col", "line", "tenure", "income", "burden", "facilities"])
    d = d[(d["tenure"] == "Renter occupied") & (d["facilities"] == "All") & (d["line"] == "Subtotal") & d["income"].isin(LOW_INCOME)]
    need = d[d["burden"] == "All"]["col"].tolist()
    burden = d[d["burden"].isin(BURDENED)]["col"].tolist()
    assert len(need) == 2 and len(burden) == 4, (need, burden)
    return need, burden


def chas(geoids: pd.Index) -> pd.DataFrame:
    need_cols, burden_cols = chas_columns()
    moe = lambda cols: [c.replace("est", "moe") for c in cols]  # noqa: E731
    with zipfile.ZipFile(CHAS_ZIP) as z:
        with z.open("140/Table8.csv") as f:
            header = pd.read_csv(f, nrows=0, encoding="latin1").columns
        gcol = next(c for c in header if c.lower() == "geoid")
        with z.open("140/Table8.csv") as f:
            t8 = pd.read_csv(f, usecols=[gcol, *need_cols, *burden_cols, *moe(need_cols), *moe(burden_cols)],
                             dtype={gcol: str}, encoding="latin1")
    t8["GEOID"] = t8[gcol].str[-11:]
    t8 = t8[t8["GEOID"].str.startswith(COUNTY_GEOID)].set_index("GEOID")
    out = pd.DataFrame(index=t8.index)
    out["need_count"] = t8[need_cols].sum(axis=1)
    out["need_count_moe"] = np.sqrt((t8[moe(need_cols)] ** 2).sum(axis=1))
    out["need_count_cv"] = cv(out["need_count"], out["need_count_moe"])
    out["burden_le50_count"] = t8[burden_cols].sum(axis=1)
    burden_moe = np.sqrt((t8[moe(burden_cols)] ** 2).sum(axis=1))
    y = out["need_count"].where(out["need_count"] > 0)
    p = (out["burden_le50_count"] / y).clip(upper=1.0)  # HUD rounding can push the part above the whole
    rad = burden_moe**2 - (p**2) * out["need_count_moe"] ** 2
    moe_p = np.where(rad >= 0, np.sqrt(rad.clip(lower=0)), np.sqrt(burden_moe**2 + (p**2) * out["need_count_moe"] ** 2)) / y
    out["chas_burden_le50_share"] = p
    out["chas_burden_le50_share_cv"] = (pd.Series(moe_p, index=out.index) / Z90) / p.where(p > 0)
    return out.reindex(geoids)


# ----------------------------------------------------------------------------------------------- SVI
def svi(geoids: pd.Index) -> pd.DataFrame:
    cols = {"RPL_THEMES": "svi_overall", "RPL_THEME1": "svi_t1", "RPL_THEME2": "svi_t2", "RPL_THEME3": "svi_t3", "RPL_THEME4": "svi_t4"}
    s = pd.read_csv(RAW / "benchmark" / "svi" / "svi2022_pa.csv", dtype={"FIPS": str}, usecols=["FIPS", *cols])
    s = s[s["FIPS"].str.startswith(COUNTY_GEOID)].set_index("FIPS").rename(columns=cols)
    s = s.mask(s < 0)  # -999 = not ranked
    return s.reindex(geoids)


# ----------------------------------------------------------------------------------------------- HCV
def hcv(geoids: pd.Index, renter_hh: pd.Series) -> pd.DataFrame:
    d = json.loads((RAW / "benchmark" / "hcv" / "hcv_by_tract_42003.json").read_text())
    rows = pd.DataFrame([f["attributes"] for f in d["features"]])
    h = rows.assign(GEOID=rows["GEOID"].astype(str)).set_index("GEOID")["HCV_PUBLIC"].astype(float)
    out = pd.DataFrame(index=geoids)
    out["hcv_count"] = h.reindex(geoids)  # absent = suppressed (<=10 holders) -> null, not zero
    out["hcv_per_renter"] = out["hcv_count"] / renter_hh.reindex(geoids).where(lambda s: s > 0)
    return out


# ----------------------------------------------------------------------------------------------- MVA
MVA21 = RAW / "benchmark" / "mva" / "mva2021.geojson"
MVA16_COUNTY = RAW / "benchmark" / "mva" / "mva2016.zip"
MVA16_CITY = RAW / "benchmark" / "mva" / "pittsmva2016.zip"


def letter_scores(letters: pd.Series) -> pd.Series:
    """Letters within one model vintage -> evenly spaced ordinal score, A = 1 ... last letter = 0. Non-letters -> NaN."""
    valid = letters.where(letters.astype(str).str.fullmatch(r"[A-J]"))
    cats = sorted(valid.dropna().unique())
    if len(cats) < 2:
        return pd.Series(np.nan, index=letters.index)
    m = {c: 1 - i / (len(cats) - 1) for i, c in enumerate(cats)}
    return valid.map(m)


def _join_letters(blocks: gpd.GeoDataFrame, polys: gpd.GeoDataFrame, col: str) -> pd.Series:
    j = gpd.sjoin(blocks[["block20", "geometry"]], polys[[col, "geometry"]].to_crs(blocks.crs), how="left", predicate="within")
    return j.drop_duplicates("block20").set_index("block20")[col]


def _tract_summary(blocks: pd.DataFrame, score: pd.Series, letter: pd.Series, prefix: str) -> pd.DataFrame:
    b = blocks.assign(score=score.reindex(blocks["block20"]).values, letter=letter.reindex(blocks["block20"]).values)
    tot = b.groupby("tract20")["hu"].sum()
    has = b[b["score"].notna()]
    wsum = has.groupby("tract20")["hu"].sum()
    mean = (has["score"] * has["hu"]).groupby(has["tract20"]).sum() / wsum.replace(0, np.nan)
    modal = has.groupby(["tract20", "letter"])["hu"].sum().reset_index().sort_values("hu", ascending=False).drop_duplicates("tract20").set_index("tract20")["letter"]
    out = pd.DataFrame({f"{prefix}_score": mean, f"{prefix}_coverage": wsum / tot.replace(0, np.nan), prefix: modal}).reindex(tot.index)
    out.loc[out[f"{prefix}_coverage"].fillna(0) < 0.5, [f"{prefix}_score", prefix]] = [np.nan, None]
    return out


def mva(geoids: pd.Index, xw: pd.DataFrame) -> pd.DataFrame:
    """Market Value Analysis 2021 and 2016 carried to 2020 tracts through block points weighted by housing units."""
    blocks = geo.blocks2020()
    blocks = blocks[blocks["tract20"].isin(geoids)].copy()
    m21 = gpd.read_file(MVA21)
    l21 = _join_letters(blocks, m21, "MVA21")
    s21 = letter_scores(l21)
    city = gpd.read_file(MVA16_CITY).rename(columns={"Cluster16": "L16"})
    county = gpd.read_file(MVA16_COUNTY).rename(columns={"MVA": "L16"})
    lc = _join_letters(blocks, city, "L16")
    lk = _join_letters(blocks, county, "L16")
    s16 = letter_scores(lc).fillna(letter_scores(lk))  # city model first (it covers Pittsburgh); county model elsewhere
    l16 = lc.where(lc.astype(str).str.fullmatch(r"[A-J]")).fillna(lk)
    out = _tract_summary(blocks, s21, l21, "mva21").join(_tract_summary(blocks, s16, l16, "mva16"))
    out["mva_change"] = out["mva21_score"] - out["mva16_score"]  # direction only
    return out.reindex(geoids)


# --------------------------------------------------------------------------------------------- flags
def flags(tracts: gpd.GeoDataFrame, xw: pd.DataFrame) -> pd.DataFrame:
    geoids = pd.Index(tracts["GEOID"])
    out = pd.DataFrame(index=geoids)

    with zipfile.ZipFile(RAW / "benchmark" / "flags" / "QCT2026CSV.zip") as z, z.open("QCT2026.csv") as f:
        q = pd.read_csv(f, dtype=str)
    out["qct"] = geoids.isin(set(q.loc[q["stcnty"] == COUNTY_GEOID, "fips"]))

    ws = openpyxl.load_workbook(RAW / "benchmark" / "flags" / "2026-DDAs-Data-Used-to-Designate.xlsx", read_only=True)["2026 MDDA"]
    rows = list(ws.iter_rows(values_only=True))
    dda = pd.DataFrame(rows[1:], columns=[str(c).strip() for c in rows[0]])
    pgh = dda[dda["Area Name"].astype(str).str.contains("Pittsburgh", case=False)]
    sdda_col = next(c for c in pgh.columns if "SDDA (1=SDDA)" in c)
    sdda = {str(int(z)).zfill(5) for z in pgh.loc[pd.to_numeric(pgh[sdda_col], errors="coerce") == 1, pgh.columns[0]]}
    zcta = gpd.read_file(RAW / "benchmark" / "flags" / "zcta2020_allegheny.geojson")[["ZCTA5", "geometry"]].to_crs(CRS_PA_SOUTH)
    ov = gpd.overlay(tracts[["GEOID", "geometry"]].to_crs(CRS_PA_SOUTH), zcta, how="intersection", keep_geom_type=False)
    ov["a"] = ov.geometry.area
    out["zcta"] = ov.sort_values("a", ascending=False).drop_duplicates("GEOID").set_index("GEOID")["ZCTA5"].reindex(geoids)
    out["dda"] = out["zcta"].isin(sdda)

    oz = json.loads((RAW / "benchmark" / "flags" / "oz_42003.json").read_text())
    oz10 = {f["attributes"]["GEOID10"] for f in oz["features"]}
    all10 = pd.Index(xw["tract10"].dropna().unique())
    ozc = geo.carry_rate(pd.Series(all10.isin(oz10).astype(float), index=all10), "tract10", xw)
    out["oz_share"] = ozc["value"].reindex(geoids)
    out["oz"] = out["oz_share"] >= 0.5
    out["xw10_dominant"] = ozc["dominant"].reindex(geoids).astype(float)

    c = pd.read_csv(RAW / "benchmark" / "flags" / "cdbg2018_pgh_bg.csv", dtype={"geoid10": str}, usecols=["geoid10", "cdbg2018"])
    cflag = pd.Series((c["cdbg2018"].astype(str).str.strip().str.lower() == "yes").astype(float).values, index=c["geoid10"])
    cflag = cflag[~cflag.index.duplicated()]
    cc = geo.carry_rate(cflag, "bg10", xw)
    out["cdbg_share"] = cc["value"].reindex(geoids)
    out["cdbg_coverage"] = cc["coverage"].reindex(geoids)
    out["cdbg"] = (out["cdbg_share"] >= 0.5).where(out["cdbg_coverage"].fillna(0) >= 0.5)
    return out


# ------------------------------------------------------------------------------------------- transit
GTFS = RAW / "benchmark" / "gtfs" / "prt_gtfs.zip"
BUFFER_M = 400.0


def transit(tracts: gpd.GeoDataFrame) -> pd.DataFrame:
    """Weekday departures at stops within 400 m of each tract, per acre of tract land."""
    with zipfile.ZipFile(GTFS) as z:
        cal = pd.read_csv(z.open("calendar.txt"), dtype={"service_id": str})
        cdates = pd.read_csv(z.open("calendar_dates.txt"), dtype={"service_id": str}) if "calendar_dates.txt" in z.namelist() else pd.DataFrame(columns=["service_id", "date", "exception_type"])
        trips = pd.read_csv(z.open("trips.txt"), usecols=["trip_id", "service_id"], dtype=str)
        stops = pd.read_csv(z.open("stops.txt"), usecols=["stop_id", "stop_lat", "stop_lon"], dtype={"stop_id": str})
        st = pd.read_csv(z.open("stop_times.txt"), usecols=["trip_id", "stop_id", "pickup_type"], dtype={"trip_id": str, "stop_id": str})
    trips_per_service = trips.groupby("service_id").size()

    def services_on(day: pd.Timestamp) -> set[str]:
        ymd = int(day.strftime("%Y%m%d"))
        col = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"][day.weekday()]
        act = set(cal.loc[(cal[col] == 1) & (cal["start_date"] <= ymd) & (cal["end_date"] >= ymd), "service_id"])
        ex = cdates[cdates["date"] == ymd]
        act |= set(ex.loc[ex["exception_type"] == 1, "service_id"])
        act -= set(ex.loc[ex["exception_type"] == 2, "service_id"])
        return act

    # representative weekday: the Wednesday in the feed's range with the most scheduled trips (regular schedule, no holiday)
    lo = pd.to_datetime(cal["start_date"].astype(str)).min()
    hi = pd.to_datetime(cal["end_date"].astype(str)).max()
    wednesdays = [d for d in pd.date_range(lo, hi, freq="D") if d.weekday() == 2]
    day = max(wednesdays, key=lambda d: trips_per_service.reindex(list(services_on(d))).fillna(0).sum())
    active = services_on(day)
    trip_ids = set(trips.loc[trips["service_id"].isin(active), "trip_id"])
    st = st[st["trip_id"].isin(trip_ids) & (st["pickup_type"].fillna(0) != 1)]
    dep = st.groupby("stop_id").size().rename("dep")
    pts = gpd.GeoDataFrame(stops, geometry=gpd.points_from_xy(stops["stop_lon"], stops["stop_lat"], crs=CRS_WGS84)).to_crs(CRS_PA_SOUTH)
    pts = pts.join(dep, on="stop_id")
    pts = pts[pts["dep"].fillna(0) > 0]
    buf = tracts[["GEOID", "ALAND", "geometry"]].to_crs(CRS_PA_SOUTH).copy()
    buf["geometry"] = buf.geometry.buffer(BUFFER_M * 3.280839895)
    j = gpd.sjoin(pts, buf, how="inner", predicate="within")
    total = j.groupby("GEOID")["dep"].sum()
    out = pd.DataFrame(index=pd.Index(tracts["GEOID"]))
    out["transit_departures"] = total.reindex(out.index).fillna(0)
    out["transit_departures_per_acre"] = out["transit_departures"] / (tracts.set_index("GEOID")["ALAND"].reindex(out.index) / SQM_PER_ACRE)
    out["transit_service_date"] = day.date().isoformat()
    print(f"  GTFS: service day {day.date()} ({len(active)} services, {len(trip_ids)} trips, {int(dep.sum()):,} departures)")
    return out


# --------------------------------------------------------------------------------------------- flood
def flood(geoids: pd.Index) -> pd.DataFrame:
    f = pd.read_csv(RAW / "hand" / "tract_ndvi_flood.csv", dtype={"GEOID": str}).set_index("GEOID")
    return f[["flood_share_pct", "flood_deep_share_pct", "veg_cover_land_pct"]].reindex(geoids)


# ------------------------------------------------------------------------------------------ eviction
ETS = RAW / "eviction" / "all_sites_monthly_2020_2021.csv"
ETS_YEARS = (2023, 2024, 2025)


def eviction(tracts: gpd.GeoDataFrame, renter_hh: pd.Series) -> pd.DataFrame:
    """Eviction Lab ETS (Pittsburgh, ZIP level) apportioned to 2020 tracts by 2020 housing units, per 100 renter HH.

    Annual filings per ZIP = mean over ETS_YEARS of the monthly `filings_2020` column (the column name is legacy; it holds
    the filings in that month). Each ZIP's filings are split among tracts in proportion to the housing units of the
    2020 blocks that fall in the ZIP's ZCTA. This is an estimate, not a tract observation, and is tagged as such.

    `eviction_zip_dominant` is the largest share of the tract's 2020 housing units that sits in one ZCTA and
    `eviction_zip_n` the number of ZCTAs that hold any of them: a tract inside one ZIP inherits that ZIP's rate cleanly,
    a tract split across several does not. (`eviction_zips` lists every ZCTA a block of the tract falls in, including
    ZCTAs that hold none of its housing, so it cannot be used for that test.)
    """
    e = pd.read_csv(ETS, dtype={"GEOID": str}, usecols=["city", "type", "GEOID", "month", "filings_2020"], low_memory=False)
    e = e[(e["city"] == "Pittsburgh, PA") & (e["type"] == "Zip Code")].copy()
    e["year"] = e["month"].str[-4:].astype(int)
    e = e[e["year"].isin(ETS_YEARS)]
    annual = e.groupby(["GEOID", "year"])["filings_2020"].sum().groupby(level=0).mean()  # mean annual filings per ZIP
    zcta = gpd.read_file(RAW / "benchmark" / "flags" / "zcta2020_allegheny.geojson")[["ZCTA5", "geometry"]].to_crs(CRS_WGS84)
    blocks = geo.blocks2020()  # whole county: a ZIP's housing outside the city still belongs to the ZIP's denominator
    j = gpd.sjoin(blocks, zcta, how="left", predicate="within").drop_duplicates("block20")
    tract_hu = j.groupby("tract20")["hu"].sum()  # every block of the tract, in a ZCTA or not
    j = j[j["ZCTA5"].notna()]
    j["zip_hu"] = j.groupby("ZCTA5")["hu"].transform("sum")
    j["zip_filings"] = j["ZCTA5"].map(annual)
    j["share"] = j["hu"] / j["zip_hu"].replace(0, np.nan)
    j["filings_part"] = j["zip_filings"] * j["share"]
    geoids = pd.Index(tracts["GEOID"])
    by_tract = j[j["tract20"].isin(geoids)]
    tot_hu = by_tract.groupby("tract20")["hu"].sum()
    covered_hu = by_tract[by_tract["zip_filings"].notna()].groupby("tract20")["hu"].sum()
    out = pd.DataFrame(index=geoids)
    out["eviction_filings_est"] = by_tract.groupby("tract20")["filings_part"].sum().reindex(geoids)
    out["eviction_coverage"] = (covered_hu / tot_hu.replace(0, np.nan)).reindex(geoids)
    out.loc[out["eviction_coverage"].fillna(0) < 0.5, "eviction_filings_est"] = np.nan
    rh = renter_hh.reindex(geoids)
    out["eviction_filing_rate"] = out["eviction_filings_est"] / rh.where(rh > 0) * 100
    out["eviction_zips"] = by_tract.groupby("tract20")["ZCTA5"].agg(lambda s: ",".join(sorted(set(s)))).reindex(geoids)
    zip_hu = by_tract.groupby(["tract20", "ZCTA5"])["hu"].sum()
    zip_hu = zip_hu[zip_hu > 0]
    out["eviction_zip_dominant"] = (zip_hu.groupby(level=0).max() / tract_hu.replace(0, np.nan)).reindex(geoids)
    out["eviction_zip_n"] = zip_hu.groupby(level=0).size().reindex(geoids).fillna(0).astype(int)
    print(f"  ETS: {annual.notna().sum()} ZIPs with filings, years {ETS_YEARS}; tracts with a rate: {out['eviction_filing_rate'].notna().sum()}")
    return out
